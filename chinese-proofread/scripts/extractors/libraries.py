"""Use already available libraries; callers own the standard-library fallback."""
import io
import re

import report
from .common import local, paragraph
from .sheets import MAX_CELLS, add_sheet, cell_text, column_index


def xlsx(raw, roots, max_chars):
    import openpyxl
    # Avoid library expansion of extreme sparse sheets. The stdlib fallback
    # reads only stored cells, retaining their actual coordinates.
    for name, root in roots.items():
        if not name.startswith('xl/worksheets/') or not name.endswith('.xml'):
            continue
        max_row, max_col = 0, 0
        for node in root.iter():
            if local(node) == 'row':
                max_row = max(max_row, int(node.get('r', '1')))
                report.require(max_row <= MAX_CELLS, '稀疏工作表改用标准库提取')
            if local(node) == 'c' and node.get('r'):
                match = re.fullmatch(r'([A-Z]{1,3})([1-9][0-9]{0,6})', node.get('r'))
                report.require(match is not None, '单元格坐标无效')
                max_col = max(max_col, column_index(match.group(1)))
                report.require(max(max_row, int(match.group(2))) * max_col <= MAX_CELLS,
                               '稀疏工作表改用标准库提取')
    workbook = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True, keep_links=False)
    blocks, total = [], 0
    try:
        for sheet in workbook.worksheets:
            sheet.reset_dimensions()
            def rows():
                nonlocal total
                for number, row in enumerate(sheet.iter_rows(), 1):
                    total += len(row)
                    report.require(number <= MAX_CELLS and total <= MAX_CELLS, '工作簿超过库解析单元格处理上限')
                    yield number, [(col, cell_text(cell.value)) for col, cell in enumerate(row, 1)]
            add_sheet(blocks, sheet.title, rows(), max_chars)
    finally:
        workbook.close()
    return blocks, ['仅提取单元格；绘图、文本框、图片和批注未提取，公式使用已保存缓存，未重新计算。'], 'openpyxl 提取全部单元格工作表（含隐藏内容）；数字和日期使用解析值，不还原显示格式'


def pptx(raw, roots, max_chars):
    from pptx import Presentation
    presentation = Presentation(io.BytesIO(raw))
    blocks, notes = [], []
    for number, slide in enumerate(presentation.slides, 1):
        counter = [0]
        def shapes(items):
            for shape in items:
                counter[0] += 1
                location = '第%d张幻灯片／对象%d' % (number, counter[0])
                if hasattr(shape, 'shapes'):
                    shapes(shape.shapes)
                elif shape.has_table:
                    for r, row in enumerate(shape.table.rows, 1):
                        for c, cell in enumerate(row.cells, 1):
                            # Spanned cells are represented by the merge origin.
                            if getattr(cell, 'is_spanned', False):
                                continue
                            paragraph(blocks, cell.text, location + '／表格 R%dC%d' % (r, c), max_chars)
                elif shape.has_text_frame:
                    for index, para in enumerate(shape.text_frame.paragraphs, 1):
                        paragraph(blocks, para.text.replace('\v', '\n'), location + '／段落%d' % index, max_chars)
                else:
                    notes.append('第%d张幻灯片存在图片、图表或其他非文本对象，未提取其文字。' % number)
        shapes(slide.shapes)
        if slide.has_notes_slide:
            frame = slide.notes_slide.notes_text_frame
            if frame is not None:
                for index, para in enumerate(frame.paragraphs, 1):
                    paragraph(blocks, para.text.replace('\v', '\n'), '第%d张幻灯片／备注段落%d' % (number, index), max_chars)
    notes.append('母版、版式继承文字、批注及备注正文之外的对象未提取；不还原视觉阅读顺序。')
    return blocks, list(dict.fromkeys(notes)), 'python-pptx 按演示顺序提取幻灯片显式文本、组合形状、表格及备注正文（含隐藏页）'
